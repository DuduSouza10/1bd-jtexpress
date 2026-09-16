from __future__ import annotations

import os
import re
import shutil
import tempfile
import zipfile
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape as xml_escape
from datetime import datetime
from pathlib import Path
from typing import Any

from flask import Flask, jsonify, render_template, request
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv("DATA_DIR", str(BASE_DIR / "data")))
DATA_DIR.mkdir(parents=True, exist_ok=True)
BACKUP_DIR = DATA_DIR / "backups"
BACKUP_DIR.mkdir(parents=True, exist_ok=True)
DATA_FILE = Path(os.getenv("DATA_FILE", str(DATA_DIR / "1B_GERAL.xlsx")))
SEED_FILE = BASE_DIR / "data" / "1B_GERAL.xlsx"
if not DATA_FILE.exists() and SEED_FILE.exists() and DATA_FILE.resolve() != SEED_FILE.resolve():
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SEED_FILE, DATA_FILE)
UPDATE_PASSWORD = os.getenv("UPDATE_PASSWORD", "3264542")
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "30"))

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024

REQUIRED = {
    "pedido": ["pedido"],
    "remessa": ["remessa"],
    "status": ["status de chamado"],
    "registro": ["data de registro"],
    "base": ["base de entrega"],
    "regional": ["regional"],
    "assinatura": ["tempo de assinatura"],
    "assinado": ["se foi assinado"],
    "rm": ["rm"],
}


def norm(text: Any) -> str:
    text = "" if text is None else str(text)
    text = text.replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", text).strip().lower()


def clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_dt(value: Any) -> datetime | None:
    if value in (None, "", 0, "0"):
        return None
    if isinstance(value, datetime):
        return value
    text = clean_value(value)
    if not text or text == "0":
        return None
    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d",
        "%d/%m/%Y",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def dt_iso(value: Any) -> str | None:
    dt = parse_dt(value)
    return dt.strftime("%Y-%m-%dT%H:%M:%S") if dt else None


def find_sheet(wb) -> Any:
    for name in wb.sheetnames:
        if norm(name) == "detalhes":
            return wb[name]
    raise ValueError("A planilha precisa conter uma aba chamada 'detalhes'.")


def header_map(ws) -> dict[str, int]:
    raw = [cell.value for cell in ws[1]]
    normalized = [norm(v) for v in raw]
    result: dict[str, int] = {}
    for key, needles in REQUIRED.items():
        if key == "rm":
            matches = [i for i, v in enumerate(normalized, start=1) if v == "rm"]
        else:
            matches = [
                i for i, v in enumerate(normalized, start=1)
                if all(needle in v for needle in needles)
            ]
        if not matches:
            raise ValueError(f"Coluna obrigatória não encontrada: {needles[0]}")
        result[key] = matches[0]
    return result



def editor_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def load_editor_page(path: Path, query: str = "", page: int = 1, per_page: int = 50) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"Arquivo de dados não encontrado: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = find_sheet(wb)
    headers = [clean_value(cell.value) or f"Coluna {idx}" for idx, cell in enumerate(ws[1], start=1)]
    max_col = len(headers)
    q = norm(query)
    matched: list[tuple[int, list[str]]] = []

    for row_number, values in enumerate(ws.iter_rows(min_row=2, max_col=max_col, values_only=True), start=2):
        display = [editor_value(v) for v in values]
        if q and q not in norm(" | ".join(display)):
            continue
        matched.append((row_number, display))

    wb.close()
    total = len(matched)
    pages = max(1, (total + per_page - 1) // per_page)
    page = max(1, min(page, pages))
    start = (page - 1) * per_page
    selected = matched[start:start + per_page]
    mtime = datetime.fromtimestamp(path.stat().st_mtime).astimezone()

    return {
        "headers": headers,
        "rows": [{"row": row_number, "values": values} for row_number, values in selected],
        "pagination": {"page": page, "perPage": per_page, "pages": pages, "total": total},
        "meta": {"updatedAt": mtime.isoformat(timespec="seconds"), "maxRow": ws.max_row, "maxCol": max_col},
    }


def details_sheet_xml_path(path: Path) -> str:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path, "r") as zin:
        workbook_root = ET.fromstring(zin.read("xl/workbook.xml"))
        sheet = None
        for candidate in workbook_root.findall(f".//{{{main_ns}}}sheet"):
            if norm(candidate.get("name")) == "detalhes":
                sheet = candidate
                break
        if sheet is None:
            raise ValueError("A planilha precisa conter uma aba chamada 'detalhes'.")
        rel_id = sheet.get(f"{{{rel_ns}}}id")
        rel_root = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        target = None
        for rel in rel_root.findall(f"{{{pkg_rel_ns}}}Relationship"):
            if rel.get("Id") == rel_id:
                target = rel.get("Target")
                break
        if not target:
            raise ValueError("Não foi possível localizar a aba 'detalhes' dentro do arquivo.")
        if target.startswith("/"):
            return target.lstrip("/")
        return f"xl/{target.lstrip('./')}"


def _build_inline_cell(ref: str, value: str, attrs: str = "") -> str:
    attrs = re.sub(r'\s+t="[^"]*"', '', attrs or "")
    attrs = re.sub(r'\s+r="[^"]*"', '', attrs)
    attrs = attrs.rstrip()
    base = f'<c r="{ref}"{attrs}>' if attrs else f'<c r="{ref}">'
    if value == "":
        return base[:-1] + "/>"
    safe = xml_escape(value.replace("\r\n", "\n").replace("\r", "\n"))
    return base[:-1] + f' t="inlineStr"><is><t xml:space="preserve">{safe}</t></is></c>'


def _replace_or_insert_cell(xml_text: str, row_number: int, col_number: int, value: str) -> str:
    ref = f"{get_column_letter(col_number)}{row_number}"
    full = re.compile(rf'<c\b(?P<attrs>[^>]*\br="{re.escape(ref)}"[^>]*)>(?P<body>.*?)</c>', re.S)
    match = full.search(xml_text)
    if match:
        replacement = _build_inline_cell(ref, value, match.group("attrs"))
        return xml_text[:match.start()] + replacement + xml_text[match.end():]

    self_closing = re.compile(rf'<c\b(?P<attrs>[^>]*\br="{re.escape(ref)}"[^>]*)/>', re.S)
    match = self_closing.search(xml_text)
    if match:
        replacement = _build_inline_cell(ref, value, match.group("attrs"))
        return xml_text[:match.start()] + replacement + xml_text[match.end():]

    if value == "":
        return xml_text

    row_re = re.compile(rf'(<row\b[^>]*\br="{row_number}"[^>]*>)(?P<body>.*?)(</row>)', re.S)
    row_match = row_re.search(xml_text)
    if not row_match:
        raise ValueError(f"Linha {row_number} não encontrada na aba detalhes.")

    body = row_match.group("body")
    new_cell = _build_inline_cell(ref, value)
    insert_at = len(body)
    for cell_match in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"[^>]*(?:/>|>.*?</c>)', body, re.S):
        existing_col = column_index_from_string(cell_match.group(1))
        if existing_col > col_number:
            insert_at = cell_match.start()
            break
    new_body = body[:insert_at] + new_cell + body[insert_at:]
    return xml_text[:row_match.start("body")] + new_body + xml_text[row_match.end("body"):]


def apply_editor_changes(path: Path, changes: list[dict[str, Any]]) -> None:
    sheet_path = details_sheet_xml_path(path)
    with zipfile.ZipFile(path, "r") as zin:
        xml_text = zin.read(sheet_path).decode("utf-8")

    for change in changes:
        row_number = int(change["row"])
        col_number = int(change["col"])
        value = "" if change.get("value") is None else str(change.get("value"))
        xml_text = _replace_or_insert_cell(xml_text, row_number, col_number, value)

    fd, tmp_name = tempfile.mkstemp(prefix="edit_1bd_", suffix=".xlsx", dir=DATA_DIR)
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp_path, "w") as zout:
            for item in zin.infolist():
                payload = xml_text.encode("utf-8") if item.filename == sheet_path else zin.read(item.filename)
                zout.writestr(item, payload)

        # Confirma que o arquivo continua íntegro e legível pelo dashboard antes de substituir.
        validated = load_data(tmp_path)
        if validated["meta"]["totalRows"] == 0:
            raise ValueError("A aba 'detalhes' não possui registros válidos após a edição.")

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(path, BACKUP_DIR / f"1B_GERAL_edit_{stamp}.xlsx")
        os.replace(tmp_path, path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


def short_status(value: str) -> str:
    return value.split("|", 1)[0].strip() if value else "Sem status"


def is_open_status(value: str) -> bool:
    s = norm(short_status(value))
    return s not in {"fechado", "processamento concluído", "processamento concluido"}


def load_data(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"Arquivo de dados não encontrado: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = find_sheet(wb)
    cols = header_map(ws)
    rows: list[dict[str, Any]] = []

    min_reg: datetime | None = None
    max_reg: datetime | None = None

    for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def cell(key: str) -> Any:
            idx = cols[key] - 1
            return row[idx] if idx < len(row) else None

        pedido = clean_value(cell("pedido"))
        remessa = clean_value(cell("remessa"))
        if not pedido and not remessa:
            continue

        status = clean_value(cell("status"))
        registro_raw = cell("registro")
        assinatura_raw = cell("assinatura")
        registro_dt = parse_dt(registro_raw)
        assinatura_dt = parse_dt(assinatura_raw)
        tempo_horas = None
        if registro_dt and assinatura_dt:
            delta = (assinatura_dt - registro_dt).total_seconds() / 3600
            if delta >= 0:
                tempo_horas = round(delta, 2)

        if registro_dt:
            min_reg = registro_dt if min_reg is None or registro_dt < min_reg else min_reg
            max_reg = registro_dt if max_reg is None or registro_dt > max_reg else max_reg

        rows.append({
            "id": r_idx,
            "pedido": pedido,
            "remessa": remessa,
            "status": status,
            "statusCurto": short_status(status),
            "registro": dt_iso(registro_raw),
            "base": clean_value(cell("base")) or "Sem base",
            "regional": clean_value(cell("regional")) or "Sem regional",
            "assinatura": dt_iso(assinatura_raw),
            "assinado": clean_value(cell("assinado")) or "Não",
            "rm": clean_value(cell("rm")) or "Sem RM",
            "tempoAssinaturaHoras": tempo_horas,
            "aberto": is_open_status(status),
        })

    wb.close()
    mtime = datetime.fromtimestamp(path.stat().st_mtime).astimezone()
    return {
        "rows": rows,
        "meta": {
            "source": path.name,
            "totalRows": len(rows),
            "updatedAt": mtime.isoformat(timespec="seconds"),
            "dateMin": min_reg.strftime("%Y-%m-%d") if min_reg else None,
            "dateMax": max_reg.strftime("%Y-%m-%d") if max_reg else None,
        },
    }


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/health")
def health():
    return jsonify({"status": "ok", "data_file": DATA_FILE.name})


@app.get("/api/data")
def api_data():
    try:
        return jsonify(load_data(DATA_FILE))
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500



@app.get("/api/editor")
def api_editor():
    try:
        query = request.args.get("q", "")[:200]
        page = max(1, int(request.args.get("page", "1")))
        per_page = int(request.args.get("per_page", "50"))
        if per_page not in {25, 50, 100, 200}:
            per_page = 50
        return jsonify(load_editor_page(DATA_FILE, query=query, page=page, per_page=per_page))
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.post("/api/editor/save")
def api_editor_save():
    payload = request.get_json(silent=True) or {}
    password = str(payload.get("password", ""))
    changes = payload.get("changes") or []

    if password != UPDATE_PASSWORD:
        return jsonify({"error": "Senha incorreta. Os dados não foram alterados."}), 403
    if not isinstance(changes, list) or not changes:
        return jsonify({"error": "Nenhuma alteração para salvar."}), 400
    if len(changes) > 5000:
        return jsonify({"error": "Muitas alterações de uma vez. Salve em blocos menores."}), 400

    try:
        wb = load_workbook(DATA_FILE, read_only=True, data_only=False)
        ws = find_sheet(wb)
        max_row, max_col = ws.max_row, ws.max_column
        wb.close()

        sanitized: list[dict[str, Any]] = []
        seen: set[tuple[int, int]] = set()
        for change in changes:
            row_number = int(change.get("row", 0))
            col_number = int(change.get("col", 0))
            if row_number < 2 or row_number > max_row:
                raise ValueError(f"Linha inválida: {row_number}")
            if col_number < 1 or col_number > max_col:
                raise ValueError(f"Coluna inválida: {col_number}")
            value = "" if change.get("value") is None else str(change.get("value"))
            if len(value) > 10000:
                raise ValueError(f"Conteúdo muito grande na célula {get_column_letter(col_number)}{row_number}.")
            key = (row_number, col_number)
            if key in seen:
                continue
            seen.add(key)
            sanitized.append({"row": row_number, "col": col_number, "value": value})

        apply_editor_changes(DATA_FILE, sanitized)
        fresh = load_data(DATA_FILE)
        return jsonify({
            "ok": True,
            "message": f"{len(sanitized)} alteração(ões) salva(s) na planilha.",
            "saved": len(sanitized),
            "meta": fresh["meta"],
        })
    except Exception as exc:
        return jsonify({"error": f"Não foi possível salvar as alterações: {exc}"}), 400


@app.post("/api/update")
def api_update():
    password = request.form.get("password", "")
    upload = request.files.get("file")

    if password != UPDATE_PASSWORD:
        return jsonify({"error": "Senha incorreta. Os dados não foram alterados."}), 403
    if not upload or not upload.filename:
        return jsonify({"error": "Selecione a planilha .xlsx antes de atualizar."}), 400
    if not upload.filename.lower().endswith(".xlsx"):
        return jsonify({"error": "Formato inválido. Envie um arquivo .xlsx."}), 400

    tmp_path: Path | None = None
    try:
        fd, tmp_name = tempfile.mkstemp(prefix="upload_1bd_", suffix=".xlsx", dir=DATA_DIR)
        os.close(fd)
        tmp_path = Path(tmp_name)
        upload.save(tmp_path)

        # Valida completamente antes de tocar no arquivo atual.
        validated = load_data(tmp_path)
        if validated["meta"]["totalRows"] == 0:
            raise ValueError("A aba 'detalhes' não possui registros válidos.")

        if DATA_FILE.exists():
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            shutil.copy2(DATA_FILE, BACKUP_DIR / f"1B_GERAL_{stamp}.xlsx")

        DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
        os.replace(tmp_path, DATA_FILE)
        tmp_path = None

        fresh = load_data(DATA_FILE)
        return jsonify({
            "ok": True,
            "message": "Dados atualizados com sucesso.",
            "meta": fresh["meta"],
        })
    except Exception as exc:
        return jsonify({"error": f"Não foi possível atualizar: {exc}"}), 400
    finally:
        if tmp_path and tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


@app.errorhandler(413)
def too_large(_):
    return jsonify({"error": f"Arquivo maior que o limite de {MAX_UPLOAD_MB} MB."}), 413


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=os.getenv("FLASK_DEBUG") == "1")
